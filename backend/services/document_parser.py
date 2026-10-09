import io
import re
import json
import hashlib
import logging
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, asdict

logger = logging.getLogger("document_parser")

COMMON_ACADEMIC_SECTIONS = [
    "Abstract",
    "Executive Summary",
    "Project Summary",
    "Specific Aims",
    "Project Description",
    "Introduction",
    "Background",
    "Methodology",
    "Research Plan",
    "Technical Approach",
    "Budget Justification",
    "Deliverables",
    "Milestones",
    "Meeting Agenda",
    "Meeting Minutes",
    "Discussion",
    "Action Items",
    "IRB Protocol",
    "Human Subjects",
    "Ethics & Compliance",
    "References",
    # Telecom & Network Engineering SOP Sections
    "Incident Summary",
    "Root Cause Analysis",
    "Problem Statement",
    "Immediate Workaround",
    "Permanent Resolution",
    "Impact Assessment",
    "Affected Cell Sites",
    "Network Topology",
    "Alarm Details",
    "Tower Clearance",
    "Safety Protocols",
    "Rollback Plan",
    "Verification Steps",
    "Action Taken"
]

@dataclass
class CanonicalDocument:
    title: str
    raw_text: str
    page_count: int
    file_type: str
    file_name: str
    file_size_bytes: int
    content_hash: str
    sections: List[Dict[str, Any]]
    metadata_hint: Dict[str, Any]

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)

def compute_sha256(content_bytes: bytes) -> str:
    """Computes SHA-256 hash of file bytes for exact deduplication."""
    return hashlib.sha256(content_bytes).hexdigest()

def detect_file_type(file_name: str, file_bytes: bytes) -> str:
    """Detects file type using magic bytes with extension fallback."""
    if file_bytes.startswith(b"%PDF"):
        return "pdf"
    if file_bytes.startswith(b"PK\x03\x04") and file_name.lower().endswith(".docx"):
        return "docx"
    
    ext = file_name.split(".")[-1].lower() if "." in file_name else ""
    if ext in ["txt", "md"]:
        return "txt"
    if ext == "json":
        return "json"
    if ext == "csv":
        return "csv"
    if ext in ["doc", "docx"]:
        return "docx"
    if ext == "pdf":
        return "pdf"
    
    return "unknown"

def extract_sections_from_text(text: str) -> List[Dict[str, Any]]:
    """Identifies standard academic & institutional sections in text."""
    sections = []
    # Build regex pattern for headers
    header_pattern = re.compile(
        r"(?:^|\n)(#{1,3}\s+|[0-9]+\.\s+)?(" + "|".join(re.escape(s) for s in COMMON_ACADEMIC_SECTIONS) + r")(?:\s*[:\-\n])",
        re.IGNORECASE
    )

    matches = list(header_pattern.finditer(text))
    if not matches:
        return [{"name": "Body", "text": text.strip(), "order": 1}]

    for idx, match in enumerate(matches):
        section_name = match.group(2).strip().title()
        start_idx = match.end()
        end_idx = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        section_content = text[start_idx:end_idx].strip()
        sections.append({
            "name": section_name,
            "text": section_content,
            "order": idx + 1
        })

    return sections

def parse_pdf(file_bytes: bytes, file_name: str) -> CanonicalDocument:
    """Parses PDF document using pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    page_count = len(reader.pages)
    extracted_pages = []

    for p in reader.pages:
        txt = p.extract_text() or ""
        extracted_pages.append(txt)

    full_text = "\n\n".join(extracted_pages).strip()

    # Determine title from first page
    title = ""
    if extracted_pages:
        lines = [line.strip() for line in extracted_pages[0].split("\n") if len(line.strip()) > 3]
        if lines:
            # First line or header that looks like a title
            title = lines[0]
            if len(title) > 250:
                title = title[:247] + "..."

    if not title:
        title = file_name.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()

    sections = extract_sections_from_text(full_text)
    content_hash = compute_sha256(file_bytes)

    return CanonicalDocument(
        title=title,
        raw_text=full_text,
        page_count=page_count,
        file_type="pdf",
        file_name=file_name,
        file_size_bytes=len(file_bytes),
        content_hash=content_hash,
        sections=sections,
        metadata_hint={}
    )

def parse_docx(file_bytes: bytes, file_name: str) -> CanonicalDocument:
    """Parses DOCX document using python-docx."""
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = []
    title = ""

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if not title and (p.style.name.startswith("Heading") or p.style.name == "Title" or len(txt) < 150):
            title = txt
        paragraphs.append(txt)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_txt:
                paragraphs.append(row_txt)

    full_text = "\n\n".join(paragraphs).strip()
    if not title:
        title = file_name.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()

    sections = extract_sections_from_text(full_text)
    content_hash = compute_sha256(file_bytes)

    return CanonicalDocument(
        title=title,
        raw_text=full_text,
        page_count=max(1, len(paragraphs) // 10),
        file_type="docx",
        file_name=file_name,
        file_size_bytes=len(file_bytes),
        content_hash=content_hash,
        sections=sections,
        metadata_hint={}
    )

def parse_txt(file_bytes: bytes, file_name: str) -> CanonicalDocument:
    """Parses plain text and markdown documents."""
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    title = lines[0].lstrip("#").strip() if lines else file_name.rsplit(".", 1)[0].title()
    sections = extract_sections_from_text(text)
    content_hash = compute_sha256(file_bytes)

    return CanonicalDocument(
        title=title,
        raw_text=text,
        page_count=max(1, len(lines) // 40),
        file_type="txt",
        file_name=file_name,
        file_size_bytes=len(file_bytes),
        content_hash=content_hash,
        sections=sections,
        metadata_hint={}
    )

def parse_json(file_bytes: bytes, file_name: str) -> CanonicalDocument:
    """Parses structured JSON grant/meeting records."""
    data = json.loads(file_bytes.decode("utf-8"))
    title = data.get("title") or data.get("project_title") or file_name.rsplit(".", 1)[0].title()
    raw_text = data.get("abstract") or data.get("content") or data.get("raw_text") or json.dumps(data, indent=2)
    content_hash = compute_sha256(file_bytes)

    sections = []
    if "sections" in data and isinstance(data["sections"], list):
        sections = data["sections"]
    else:
        sections = extract_sections_from_text(raw_text)

    return CanonicalDocument(
        title=title,
        raw_text=raw_text,
        page_count=1,
        file_type="json",
        file_name=file_name,
        file_size_bytes=len(file_bytes),
        content_hash=content_hash,
        sections=sections,
        metadata_hint=data
    )

def parse_document(file_name: str, file_bytes: bytes) -> CanonicalDocument:
    """
    Main entry point for document parsing.
    Validates file type and extracts CanonicalDocument representation.
    """
    detected_type = detect_file_type(file_name, file_bytes)

    if detected_type == "pdf":
        return parse_pdf(file_bytes, file_name)
    elif detected_type == "docx":
        return parse_docx(file_bytes, file_name)
    elif detected_type in ["txt", "md", "csv"]:
        return parse_txt(file_bytes, file_name)
    elif detected_type == "json":
        return parse_json(file_bytes, file_name)
    else:
        raise ValueError(f"CAP-1001: Unsupported file type for '{file_name}'. Allowed: PDF, DOCX, TXT, JSON, CSV.")
